"""Review CSV → columns in source/Wloski.xlsx (example sentences and/or topics)

Usage (from the project folder):
    python tools/import_examples.py source/przyklady_A1.csv
    python tools/import_examples.py source/tematy.csv

The CSV (semicolon-separated, UTF-8) needs columns Włoski, Type, Rodzaj and at least one of:
Przykład, Przykład PL, Przykład EN, Temat. Only the columns present in the CSV are imported.
Rows are matched to the sheet "Słownik" by Włoski + Type + Rodzaj (the same key as the word id).
Missing columns are added to the sheet. Existing values are overwritten.

    python tools/import_examples.py --add source/czasowniki_zwrotne.csv
With --add, CSV rows not found in the sheet are appended as new words (every CSV column that the sheet has,
so the CSV should have all columns of "Słownik": Polski, Angielski, Typ, Tipo, Poziom...).
Afterwards run: python tools/excel_to_json.py (for new verbs first: python tools/conjugate.py)
"""
import csv
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Brak biblioteki openpyxl. Zainstaluj: pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "source" / "Wloski.xlsx"
IMPORTABLE = ["Przykład", "Przykład PL", "Przykład EN", "Temat"]


def key(it, pos, g):
    return (str(it or "").strip(), str(pos or "").strip().lower(), str(g or "").strip().lower())


def main():
    args = sys.argv[1:]
    add = "--add" in args
    args = [a for a in args if a != "--add"]
    if len(args) != 1:
        sys.exit(__doc__)
    with open(args[0], encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f, delimiter=";"))
    missing = [c for c in ["Włoski", "Type", "Rodzaj"] if rows and c not in rows[0]]
    if missing:
        sys.exit(f"Brak kolumn w CSV: {', '.join(missing)}")
    EXAMPLE = [c for c in IMPORTABLE if rows and c in rows[0]]
    if not EXAMPLE:
        sys.exit(f"CSV nie ma żadnej z kolumn: {', '.join(IMPORTABLE)}")
    new = {key(r["Włoski"], r["Type"], r["Rodzaj"]): [r[c].strip() for c in EXAMPLE] for r in rows}

    wb = openpyxl.load_workbook(SRC)
    ws = wb["Słownik"]
    head = [c.value for c in ws[1]]
    for name in EXAMPLE:
        if name not in head:
            ws.cell(row=1, column=len(head) + 1, value=name)
            head.append(name)
    col = {n: head.index(n) + 1 for n in ["Włoski", "Type", "Rodzaj"] + EXAMPLE}

    done, changed = set(), 0
    for n in range(2, ws.max_row + 1):
        k = key(*(ws.cell(row=n, column=col[c]).value for c in ["Włoski", "Type", "Rodzaj"]))
        if k not in new:
            continue
        for name, value in zip(EXAMPLE, new[k]):
            cell = ws.cell(row=n, column=col[name])
            if (cell.value or "") != value:
                cell.value = value or None
                changed += 1
        done.add(k)

    lost = [k for k in new if k not in done]
    if add and lost:
        need = [c for c in ["Polski", "Angielski", "Poziom"] if c not in rows[0]]
        if need:
            sys.exit(f"--add: brak kolumn w CSV: {', '.join(need)}")
        for r in rows:
            if key(r["Włoski"], r["Type"], r["Rodzaj"]) in lost:
                ws.append([(r.get(h) or "").strip() or None if h in r else None for h in head])
                changed += 1
        print(f"Dopisano {len(lost)} nowych słów: " + ", ".join(k[0] for k in lost))
        done.update(lost)
        lost = []
        if ws.auto_filter.ref:
            ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(head))}{ws.max_row}"
    if lost:
        print("Nie znaleziono w Excelu (pominięte):")
        print("\n".join(f"  - {k[0]} ({k[1]} {k[2]})".rstrip() for k in lost))
    wb.save(SRC)
    print(f"Wgrano {', '.join(EXAMPLE)} dla {len(done)} słów ({changed} zmienionych komórek) do {SRC.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
