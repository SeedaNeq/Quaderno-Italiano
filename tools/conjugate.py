"""Generator odmiany czasowników → arkusz "Odmiana" w source/Wloski.xlsx

Usage (from the project folder):
    python tools/conjugate.py              dopisuje odmianę czasowników, których jeszcze nie ma w arkuszu
    python tools/conjugate.py andare bere  generuje od nowa podane czasowniki (nadpisuje ręczne poprawki!)
    python tools/conjugate.py --print bere pokazuje odmianę bez zapisu

Formy powstają z reguł (-are/-ere/-ire, -isc-, pisownia -care/-gare/-ciare/-giare) i z tabeli IRR
z czasownikami nieregularnymi (też z przedrostkami: tradurre, proporre, mantenere...).
Arkusz "Odmiana" to źródło prawdy – poprawki robisz w Excelu, ten skrypt ich nie rusza (poza jawnie podanymi).
Kolumny: Czasownik · Czas · io · tu · lui/lei · noi · voi · loro. Kilka wariantów rozdziel ";".
Przy czasach złożonych z essere: "sono andato/a", "siamo andati/e". Imperativo: io puste, lui/lei = Lei, loro = Loro.
"""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "source" / "Wloski.xlsx"
SHEET = "Odmiana"

# key, Italian name (Excel column "Czas"), level. Keep in sync with TENSES in app/js/verbs.js.
TENSES = [
    ("pres", "presente", "A1"),
    ("pp", "passato prossimo", "A2"), ("impf", "imperfetto", "A2"), ("fut", "futuro semplice", "A2"), ("imp", "imperativo", "A2"),
    ("cond", "condizionale presente", "B1"), ("cong", "congiuntivo presente", "B1"), ("trap", "trapassato prossimo", "B1"),
    ("congi", "congiuntivo imperfetto", "B2"), ("congp", "congiuntivo passato", "B2"), ("condp", "condizionale passato", "B2"),
    ("rem", "passato remoto", "B2"),
]
PERSONS = ["io", "tu", "lui/lei", "noi", "voi", "loro"]

# -ire verbs WITHOUT -isc- (all other -ire verbs get it: capisco, finisco...)
NO_ISC = {"aprire", "dormire", "partire", "sentire", "offrire", "servire", "seguire", "vestire", "avvertire", "coprire",
          "divertire", "fuggire", "mentire", "scoprire", "soffrire", "investire", "consentire", "eseguire", "proseguire",
          "perseguire", "invertire", "risentire", "sfuggire", "travestire", "bollire", "pentire"}
# -iare verbs with stressed i: invio → tu invii
STRESSED_I = {"inviare", "avviare", "rinviare", "deviare", "ovviare", "sciare", "spiare"}
# auxiliary essere (reflexive verbs get it automatically); BOTH = both forms are given
ESSERE = {"andare", "arrivare", "entrare", "partire", "tornare", "uscire", "venire", "essere", "stare", "piacere", "costare",
          "diventare", "cadere", "morire", "nascere", "restare", "rimanere", "riuscire", "ritornare", "sembrare", "succedere",
          "bastare", "dispiacere", "mancare", "crescere", "dimagrire", "ingrassare", "scappare", "scivolare", "accadere",
          "apparire", "appartenere", "bisognare", "capitare", "comparire", "dipendere", "esistere", "fuggire", "intervenire",
          "provenire", "rientrare", "risultare", "scomparire", "scoppiare", "sparire", "sopravvivere", "importare", "divenire",
          "emergere", "emigrare", "immigrare", "occorrere", "parere", "prevalere", "precipitare", "scadere", "sfociare",
          "sfuggire", "subentrare", "spettare", "giungere", "crollare", "calare", "derivare", "arrossire", "valere",
          "equivalere", "affogare", "dilagare", "ricorrere", "scorrere"}
BOTH = {"salire", "scendere", "passare", "vivere", "durare", "risalire"}
NO_IMP = {"dovere", "potere", "volere", "bisognare", "piacere", "dispiacere", "parere", "occorrere", "accadere", "succedere",
          "capitare", "bastare", "sembrare", "costare", "spettare", "importare", "esistere", "dipendere", "risultare",
          "equivalere", "vertere", "competere"}
THIRD_ONLY = {"bisognare", "accadere", "succedere", "occorrere", "spettare"}   # only lui/lei (+ loro except bisognare)
NO_PP = {"vertere", "competere", "bisognare"}                                      # no participle → no compound tenses

# Irregular verbs. Fields (all optional):
#   long  – "long" infinitive used for imperfetto, congiuntivo imperfetto, weak passato remoto (fare → facere)
#   pres, cong, congi, impf, imp, remf – full lists of 6 forms; fut – future stem; rem – strong remoto stem
#   (misi, mise, misero + weak mettesti, mettemmo, metteste); pp – participle
IRR = {
    "essere": dict(pres=["sono", "sei", "è", "siamo", "siete", "sono"], impf=["ero", "eri", "era", "eravamo", "eravate", "erano"],
                   fut="sar", cong=["sia", "sia", "sia", "siamo", "siate", "siano"],
                   congi=["fossi", "fossi", "fosse", "fossimo", "foste", "fossero"], remf=["fui", "fosti", "fu", "fummo", "foste", "furono"],
                   pp="stato", imp=["", "sii", "sia", "siamo", "siate", "siano"]),
    "avere": dict(pres=["ho", "hai", "ha", "abbiamo", "avete", "hanno"], fut="avr", cong=["abbia", "abbia", "abbia", "abbiamo", "abbiate", "abbiano"],
                  rem="ebb", imp=["", "abbi", "abbia", "abbiamo", "abbiate", "abbiano"]),
    "andare": dict(pres=["vado", "vai", "va", "andiamo", "andate", "vanno"], fut="andr", imp=["", "va'; vai", "vada", "andiamo", "andate", "vadano"]),
    "dare": dict(pres=["do", "dai", "dà", "diamo", "date", "danno"], fut="dar", cong=["dia", "dia", "dia", "diamo", "diate", "diano"],
                 congi=["dessi", "dessi", "desse", "dessimo", "deste", "dessero"],
                 remf=["diedi; detti", "desti", "diede; dette", "demmo", "deste", "diedero; dettero"], imp=["", "da'; dai", "dia", "diamo", "date", "diano"]),
    "stare": dict(pres=["sto", "stai", "sta", "stiamo", "state", "stanno"], fut="star", cong=["stia", "stia", "stia", "stiamo", "stiate", "stiano"],
                  congi=["stessi", "stessi", "stesse", "stessimo", "steste", "stessero"],
                  remf=["stetti", "stesti", "stette", "stemmo", "steste", "stettero"], imp=["", "sta'; stai", "stia", "stiamo", "state", "stiano"]),
    "fare": dict(long="facere", pres=["faccio", "fai", "fa", "facciamo", "fate", "fanno"], fut="far", rem="fec", pp="fatto",
                 imp=["", "fa'; fai", "faccia", "facciamo", "fate", "facciano"]),
    "soddisfare": dict(long="soddisfacere", pres=["soddisfo", "soddisfi", "soddisfa", "soddisfiamo", "soddisfate", "soddisfano"],
                       cong=["soddisfi", "soddisfi", "soddisfi", "soddisfiamo", "soddisfiate", "soddisfino"], fut="soddisfer",
                       rem="soddisfec", pp="soddisfatto", imp=["", "soddisfa", "soddisfi", "soddisfiamo", "soddisfate", "soddisfino"]),
    "dire": dict(long="dicere", pres=["dico", "dici", "dice", "diciamo", "dite", "dicono"], fut="dir", rem="diss", pp="detto",
                 imp=["", "di'", "dica", "diciamo", "dite", "dicano"]),
    "sapere": dict(pres=["so", "sai", "sa", "sappiamo", "sapete", "sanno"], fut="sapr", cong=["sappia", "sappia", "sappia", "sappiamo", "sappiate", "sappiano"],
                   rem="sepp", imp=["", "sappi", "sappia", "sappiamo", "sappiate", "sappiano"]),
    "dovere": dict(pres=["devo", "devi", "deve", "dobbiamo", "dovete", "devono"], fut="dovr", cong=["debba", "debba", "debba", "dobbiamo", "dobbiate", "debbano"]),
    "potere": dict(pres=["posso", "puoi", "può", "possiamo", "potete", "possono"], fut="potr"),
    "volere": dict(pres=["voglio", "vuoi", "vuole", "vogliamo", "volete", "vogliono"], fut="vorr", rem="voll"),
    "venire": dict(pres=["vengo", "vieni", "viene", "veniamo", "venite", "vengono"], fut="verr", rem="venn", pp="venuto"),
    "tenere": dict(pres=["tengo", "tieni", "tiene", "teniamo", "tenete", "tengono"], fut="terr", rem="tenn"),
    "rimanere": dict(pres=["rimango", "rimani", "rimane", "rimaniamo", "rimanete", "rimangono"], fut="rimarr", rem="rimas", pp="rimasto"),
    "uscire": dict(pres=["esco", "esci", "esce", "usciamo", "uscite", "escono"]),
    "bere": dict(long="bevere", fut="berr", rem="bevv"),
    "morire": dict(pres=["muoio", "muori", "muore", "moriamo", "morite", "muoiono"], pp="morto"),
    "salire": dict(pres=["salgo", "sali", "sale", "saliamo", "salite", "salgono"]),
    "parire": dict(pres=["paio", "pari", "pare", "pariamo", "parite", "paiono"], rem="parv", pp="parso"),   # apparire, comparire
    "porre": dict(long="ponere", pres=["pongo", "poni", "pone", "poniamo", "ponete", "pongono"], fut="porr", rem="pos", pp="posto"),
    "durre": dict(long="ducere", fut="durr", rem="duss", pp="dotto"),                                        # tradurre, produrre...
    "trarre": dict(long="traere", pres=["traggo", "trai", "trae", "traiamo", "traete", "traggono"], fut="trarr", rem="trass", pp="tratto"),
    "valere": dict(pres=["valgo", "vali", "vale", "valiamo", "valete", "valgono"], fut="varr", rem="vals", pp="valso"),
    "parere": dict(pres=["paio", "pari", "pare", "paiamo", "parete", "paiono"], fut="parr", rem="parv", pp="parso"),
    "piacere": dict(pres=["piaccio", "piaci", "piace", "piacciamo", "piacete", "piacciono"], rem="piacqu", pp="piaciuto"),
    "tacere": dict(pres=["taccio", "taci", "tace", "tacciamo", "tacete", "tacciono"], rem="tacqu", pp="taciuto"),
    "nuocere": dict(pres=["noccio", "nuoci", "nuoce", "nociamo", "nocete", "nocciono"], rem="nocqu", pp="nociuto"),
    "sedere": dict(pres=["siedo", "siedi", "siede", "sediamo", "sedete", "siedono"]),
    "cadere": dict(fut="cadr", rem="cadd"),
    "vedere": dict(fut="vedr", rem="vid", pp="visto"),
    "provvedere": dict(rem="provvid"),
    "godere": dict(fut="godr"),
    "vivere": dict(fut="vivr", rem="viss", pp="vissuto"),
    "scegliere": dict(pres=["scelgo", "scegli", "sceglie", "scegliamo", "scegliete", "scelgono"], rem="scels", pp="scelto"),
    "togliere": dict(pres=["tolgo", "togli", "toglie", "togliamo", "togliete", "tolgono"], rem="tols", pp="tolto"),
    "cogliere": dict(pres=["colgo", "cogli", "coglie", "cogliamo", "cogliete", "colgono"], rem="cols", pp="colto"),
    "spegnere": dict(pres=["spengo", "spegni", "spegne", "spegniamo", "spegnete", "spengono"], rem="spens", pp="spento"),
    "riempire": dict(pres=["riempio", "riempi", "riempie", "riempiamo", "riempite", "riempiono"]),
    "cucire": dict(pres=["cucio", "cuci", "cuce", "cuciamo", "cucite", "cuciono"]),
    "compiere": dict(pres=["compio", "compi", "compie", "compiamo", "compiete", "compiono"], pp="compiuto"),
    "conoscere": dict(rem="conobb", pp="conosciuto"),
    "crescere": dict(rem="crebb", pp="cresciuto"),
    "nascere": dict(rem="nacqu", pp="nato"),
    "mettere": dict(rem="mis", pp="messo"),
    "prendere": dict(rem="pres", pp="preso"),
    "scrivere": dict(rem="scriss", pp="scritto"),
    "correre": dict(rem="cors", pp="corso"),
    "muovere": dict(rem="moss", pp="mosso"),
    "rompere": dict(rem="rupp", pp="rotto"),
    "perdere": dict(rem="pers", pp="perso"),
    "leggere": dict(rem="less", pp="letto"),
    "correggere": dict(rem="corress", pp="corretto"),
    "proteggere": dict(rem="protess", pp="protetto"),
    "reggere": dict(rem="ress", pp="retto"),
    "dirigere": dict(rem="diress", pp="diretto"),
    "redigere": dict(rem="redass", pp="redatto"),
    "distruggere": dict(rem="distruss", pp="distrutto"),
    "friggere": dict(rem="friss", pp="fritto"),
    "volgere": dict(rem="vols", pp="volto"),
    "chiudere": dict(rem="chius", pp="chiuso"),
    "cludere": dict(rem="clus", pp="cluso"),
    "ludere": dict(rem="lus", pp="luso"),
    "rispondere": dict(rem="rispos", pp="risposto"),
    "nascondere": dict(rem="nascos", pp="nascosto"),
    "chiedere": dict(rem="chies", pp="chiesto"),
    "cidere": dict(rem="cis", pp="ciso"),
    "ridere": dict(rem="ris", pp="riso"),
    "dividere": dict(rem="divis", pp="diviso"),
    "giungere": dict(rem="giuns", pp="giunto"),
    "stringere": dict(rem="strins", pp="stretto"),
    "spingere": dict(rem="spins", pp="spinto"),
    "dipingere": dict(rem="dipins", pp="dipinto"),
    "fingere": dict(rem="fins", pp="finto"),
    "piangere": dict(rem="pians", pp="pianto"),
    "vincere": dict(rem="vins", pp="vinto"),
    "distinguere": dict(rem="distins", pp="distinto"),
    "infrangere": dict(rem="infrans", pp="infranto"),
    "accorgere": dict(rem="accors", pp="accorto"),
    "emergere": dict(rem="emers", pp="emerso"),
    "mordere": dict(rem="mors", pp="morso"),
    "discutere": dict(rem="discuss", pp="discusso"),
    "risolvere": dict(rem="risols", pp="risolto"),
    "suadere": dict(rem="suas", pp="suaso"),
    "fondere": dict(rem="fus", pp="fuso"),
    "sumere": dict(rem="suns", pp="sunto"),
    "primere": dict(rem="press", pp="presso"),
    "concedere": dict(rem="conces", pp="concesso"),
    "succedere": dict(rem="success", pp="successo"),
    "espandere": dict(rem="espans", pp="espanso"),
    "evolvere": dict(pp="evoluto"),
    "devolvere": dict(pp="devoluto"),
    "esigere": dict(pp="esatto"),
    "aprire": dict(pp="aperto"),
    "coprire": dict(pp="coperto"),
    "frire": dict(pp="ferto"),                                                                               # offrire, soffrire
    "istere": dict(pp="istito"),                                                                             # esistere, insistere...
}
# -endere verbs: accendere → acceso, accesi
for v in ["accendere", "appendere", "attendere", "difendere", "dipendere", "scendere", "spendere", "rendere", "intendere",
          "tendere", "offendere", "arrendere"]:
    IRR[v] = dict(rem=v[:-5] + "s", pp=v[:-5] + "so")
# Verbs built on an irregular base: verb → base (the verb is prefix + base)
DERIVED = {
    "rifare": "fare", "contraddire": "dire", "disdire": "dire",
    "intervenire": "venire", "provenire": "venire", "divenire": "venire",
    "appartenere": "tenere", "contenere": "tenere", "mantenere": "tenere", "ottenere": "tenere", "sostenere": "tenere",
    "trattenere": "tenere", "intrattenere": "tenere", "ritenere": "tenere", "attenere": "tenere",
    "comporre": "porre", "disporre": "porre", "esporre": "porre", "imporre": "porre", "opporre": "porre",
    "predisporre": "porre", "proporre": "porre", "sottoporre": "porre",
    "condurre": "durre", "introdurre": "durre", "produrre": "durre", "ridurre": "durre", "tradurre": "durre",
    "distrarre": "trarre", "riuscire": "uscire", "risalire": "salire",
    "apparire": "parire", "comparire": "parire", "scomparire": "parire",
    "equivalere": "valere", "prevalere": "valere", "avvalere": "valere", "dispiacere": "piacere", "possedere": "sedere",
    "accadere": "cadere", "scadere": "cadere", "prevedere": "vedere", "rivedere": "vedere",
    "convivere": "vivere", "sopravvivere": "vivere", "accogliere": "cogliere", "raccogliere": "cogliere",
    "riconoscere": "conoscere", "commuovere": "muovere", "promuovere": "muovere",
    "ammettere": "mettere", "commettere": "mettere", "compromettere": "mettere", "permettere": "mettere",
    "promettere": "mettere", "rimettere": "mettere", "scommettere": "mettere", "smettere": "mettere",
    "trasmettere": "mettere", "dimettere": "mettere",
    "comprendere": "prendere", "riprendere": "prendere", "intraprendere": "prendere",
    "descrivere": "scrivere", "iscrivere": "scrivere", "sottoscrivere": "scrivere",
    "ricorrere": "correre", "scorrere": "correre", "trascorrere": "correre", "occorrere": "correre",
    "interrompere": "rompere", "avvolgere": "volgere", "rivolgere": "volgere", "svolgere": "volgere",
    "coinvolgere": "volgere", "capovolgere": "volgere",
    "concludere": "cludere", "escludere": "cludere", "precludere": "cludere",
    "alludere": "ludere", "deludere": "ludere", "eludere": "ludere", "illudere": "ludere",
    "corrispondere": "rispondere", "richiedere": "chiedere",
    "decidere": "cidere", "uccidere": "cidere", "incidere": "cidere", "sorridere": "ridere", "condividere": "dividere",
    "aggiungere": "giungere", "raggiungere": "giungere", "costringere": "stringere", "convincere": "vincere",
    "soffriggere": "friggere", "persuadere": "suadere", "dissuadere": "suadere",
    "confondere": "fondere", "diffondere": "fondere",
    "assumere": "sumere", "desumere": "sumere", "riassumere": "sumere",
    "esprimere": "primere", "deprimere": "primere", "reprimere": "primere", "sopprimere": "primere",
    "scoprire": "coprire", "offrire": "frire", "soffrire": "frire",
    "esistere": "istere", "insistere": "istere", "resistere": "istere", "assistere": "istere",
    "pretendere": "tendere", "stendere": "tendere",
}
REFLEXIVE_BASE = {"distrarsi": "distrarre", "disporsi": "disporre", "proporsi": "proporre"}

END = {
    "pres": {"are": ["o", "i", "a", "iamo", "ate", "ano"], "ere": ["o", "i", "e", "iamo", "ete", "ono"], "ire": ["o", "i", "e", "iamo", "ite", "ono"]},
    "fut": ["ò", "ai", "à", "emo", "ete", "anno"],
    "cond": ["ei", "esti", "ebbe", "emmo", "este", "ebbero"],
    "impf": ["vo", "vi", "va", "vamo", "vate", "vano"],
    "congi": ["ssi", "ssi", "sse", "ssimo", "ste", "ssero"],
}
ACC = {"a": "à", "e": "è", "i": "ì", "o": "ò", "u": "ù"}


def variants(s):
    return [x.strip() for x in s.split(";")] if s else []


def each(s, f):
    """Apply f to every variant of a form like "va'; vai"."""
    return "; ".join(f(x) for x in variants(s)) if s else ""


def regular(inf):
    """Simple tenses + participle of a verb by the regular rules (inf = infinitive or long infinitive)."""
    c, st = inf[-3:], inf[:-3]
    f = {}
    if c == "are":
        hard = st + "h" if st.endswith(("c", "g")) else st                     # cerc → cerchi, cerchiamo
        if st.endswith("i"):
            tu, noi = (st + "i" if inf in STRESSED_I else st), st + "amo"
        else:
            tu, noi = hard + "i", hard + "iamo"
        f["pres"] = [st + "o", tu, st + "a", noi, st + "ate", st + "ano"]
        f["cong"] = [tu, tu, tu, noi, noi[:-3] + "ate", tu + "no"]
        fut = (st[:-1] if st.endswith(("ci", "gi")) else hard) + "er"
        f["rem"] = [st + "ai", st + "asti", st + "ò", st + "ammo", st + "aste", st + "arono"]
        f["pp"] = st + "ato"
        imp_tu = st + "a"
    else:
        isc = c == "ire" and inf not in NO_ISC
        if isc:
            f["pres"] = [st + "isco", st + "isci", st + "isce", st + "iamo", st + "ite", st + "iscono"]
        else:
            f["pres"] = [st + e for e in END["pres"][c]]
        io = f["pres"][0][:-1]
        f["cong"] = [io + "a", io + "a", io + "a", f["pres"][3], f["pres"][3][:-3] + "ate", io + "ano"]
        fut = st + c[0] + "r"
        if c == "ere":
            two = not st.endswith("t")                                          # vendetti, but not ripetetti
            f["rem"] = [st + "ei" + ("; " + st + "etti" if two else ""), st + "esti", st + "é" + ("; " + st + "ette" if two else ""),
                        st + "emmo", st + "este", st + "erono" + ("; " + st + "ettero" if two else "")]
            f["pp"] = st + ("iuto" if st.endswith("c") else "uto")
        else:
            f["rem"] = [st + "ii", st + "isti", st + "ì", st + "immo", st + "iste", st + "irono"]
            f["pp"] = st + "ito"
        imp_tu = f["pres"][1]
    f["fut"] = [fut + e for e in END["fut"]]
    f["cond"] = [fut + e for e in END["cond"]]
    f["impf"] = [inf[:-2] + e for e in END["impf"]]
    f["congi"] = [inf[:-2] + e for e in END["congi"]]
    f["imp"] = ["", imp_tu, f["cong"][2], f["pres"][3], f["pres"][4], f["cong"][5]]
    return f


def irregular(inf, spec, prefix=""):
    """Regular forms of the (long) infinitive with the irregular parts of spec laid over them."""
    P = lambda s: each(s, lambda x: join_prefix(prefix, x))
    long_inf = prefix + spec.get("long", inf[len(prefix):])
    f = regular(long_inf)
    if "pres" in spec:
        pres = [P(x) for x in spec["pres"]]
        f["pres"] = pres
        io = pres[0][:-1]
        f["cong"] = [io + "a", io + "a", io + "a", pres[3], pres[3][:-3] + "ate", io + "ano"]
        f["imp"] = ["", pres[1] if not long_inf.endswith("are") else f["imp"][1], f["cong"][2], pres[3], pres[4], f["cong"][5]]
    if "cong" in spec:
        f["cong"] = [P(x) for x in spec["cong"]]
        f["imp"][2], f["imp"][5] = f["cong"][2], f["cong"][5]
    if "fut" in spec:
        fut = prefix + spec["fut"]
        f["fut"] = [fut + e for e in END["fut"]]
        f["cond"] = [fut + e for e in END["cond"]]
    if "rem" in spec:
        s = prefix + spec["rem"]
        f["rem"] = [s + "i", f["rem"][1], s + "e", f["rem"][3], f["rem"][4], s + "ero"]
    for k in ("impf", "congi", "imp"):
        if k in spec:
            f[k] = [P(x) for x in spec[k]]
    if "remf" in spec:
        f["rem"] = [P(x) for x in spec["remf"]]
    if "pp" in spec:
        f["pp"] = prefix + spec["pp"]
    if prefix and inf.endswith("dire"):
        f["imp"][1] = prefix + "dici"                                          # contraddici, not contraddi'
    return f


def join_prefix(prefix, form):
    if not prefix or not form:
        return form
    if form.endswith("'"):                                                      # rifare: fa' → rifà
        form = form[:-2] + ACC[form[-2]]
    elif len(form) <= 3 and form[-1] in ACC and form in ("fa", "va", "da", "sta"):
        form = form[:-1] + ACC[form[-1]]
    return prefix + form


def simple(inf):
    if inf in IRR:
        return irregular(inf, IRR[inf])
    if inf in DERIVED:
        base = DERIVED[inf]
        return irregular(inf, IRR[base], inf[:-len(base)])
    return regular(inf)


AUX = {
    "avere": {"pres": ["ho", "hai", "ha", "abbiamo", "avete", "hanno"], "impf": ["avevo", "avevi", "aveva", "avevamo", "avevate", "avevano"],
              "cong": ["abbia", "abbia", "abbia", "abbiamo", "abbiate", "abbiano"], "cond": ["avrei", "avresti", "avrebbe", "avremmo", "avreste", "avrebbero"]},
    "essere": {"pres": ["sono", "sei", "è", "siamo", "siete", "sono"], "impf": ["ero", "eri", "era", "eravamo", "eravate", "erano"],
               "cong": ["sia", "sia", "sia", "siamo", "siate", "siano"], "cond": ["sarei", "saresti", "sarebbe", "saremmo", "sareste", "sarebbero"]},
}
COMPOUND = {"pp": "pres", "trap": "impf", "congp": "cong", "condp": "cond"}


def agree(pp, i, fem=False):
    if fem:
        return pp[:-1] + "a"
    return pp[:-1] + ("i/e" if i >= 3 else "o/a")


def clitic(pron, form):
    """mi + alzo → mi alzo; se la + è → se l'è"""
    if pron.endswith((" la", " ne")) and form[:1] in "aeiouèé":
        return pron[:-1] + "'" + form
    return pron + " " + form


def conjugate(verb):
    """All tenses of a dictionary verb (may be reflexive and/or several words): {tense key: [6 forms]}."""
    head, _, tail = verb.partition(" ")
    tail = " " + tail if tail else ""
    kind = ""
    for k in ("sela", "sene", "si"):
        if head.endswith("r" + k):
            kind = k
            break
    inf = REFLEXIVE_BASE.get(head) or (head[:-len(kind)] + ("re" if head[:-len(kind)].endswith("or") else "e") if kind else head)
    f = simple(inf)
    aux = ["essere"] if kind or inf in ESSERE else ["avere", "essere"] if inf in BOTH else ["avere"]

    if kind:
        pr = {"si": ["mi", "ti", "si", "ci", "vi", "si"], "sela": ["me la", "te la", "se la", "ce la", "ve la", "se la"],
              "sene": ["me ne", "te ne", "se ne", "ce ne", "ve ne", "se ne"]}[kind]
        enc = {"si": ["", "ti", "", "ci", "vi", ""], "sela": ["", "tela", "", "cela", "vela", ""], "sene": ["", "tene", "", "cene", "vene", ""]}[kind]
        for t in ("pres", "impf", "fut", "cond", "cong", "congi", "rem"):
            f[t] = [each(x, lambda y, i=i: clitic(pr[i], y)) for i, x in enumerate(f[t])]
        f["imp"] = [each(x, lambda y, i=i: (y + enc[i]) if enc[i] else clitic(pr[i], y)) if x else "" for i, x in enumerate(f["imp"])]
        tu = variants(f["imp"][1])
        short = [x for x in tu if "'" in x]                        # fa' + ti → fatti (fa'ti / faiti do not exist)
        if short:
            f["imp"][1] = "; ".join(x.replace("'", x[x.index("'") + 1]) for x in short)
    pp = f.pop("pp")
    for t, base in COMPOUND.items():
        out = []
        for i in range(6):
            forms = []
            for a in aux:
                part = agree(pp, i, kind == "sela") if a == "essere" else pp
                s = AUX[a][base][i] + " " + part
                forms.append(clitic(pr[i], s) if kind else s)
            out.append("; ".join(forms))
        f[t] = out
    if inf in NO_PP:
        for t in COMPOUND:
            f[t] = [""] * 6
    if inf in NO_IMP:
        f["imp"] = [""] * 6
    if inf in THIRD_ONLY:
        keep = {2} if inf == "bisognare" else {2, 5}
        f = {t: [x if i in keep else "" for i, x in enumerate(v)] for t, v in f.items()}
        f["imp"] = [""] * 6
    if tail:
        f = {t: [each(x, lambda y: y + tail) for x in v] for t, v in f.items()}
    return {t: f[t] for t, _, _ in TENSES}


def regular_only(verb):
    """Same verb conjugated with no irregular table – used to mark verbs as irregular."""
    saved = dict(IRR), dict(DERIVED)
    IRR.clear(); DERIVED.clear()
    try:
        return conjugate(verb)
    finally:
        IRR.update(saved[0]); DERIVED.update(saved[1])


def main(args):
    import openpyxl
    if args and args[0] == "--print":
        for v in args[1:]:
            for (k, name, _), forms in zip(TENSES, conjugate(v).values()):
                print(f"{v:14} {name:24} " + " | ".join(forms))
        return
    wb = openpyxl.load_workbook(SRC)
    words = wb["Słownik"]
    head = [c.value for c in words[1]]
    verbs = sorted({r[head.index("Włoski")].strip() for r in words.iter_rows(min_row=2, values_only=True)
                    if r[head.index("Type")] and str(r[head.index("Type")]).strip().lower() == "verb" and r[head.index("Włoski")]})
    if SHEET in wb.sheetnames:
        ws = wb[SHEET]
        old = [list(r) for r in ws.iter_rows(min_row=2, values_only=True) if r and r[0]]
    else:
        ws = wb.create_sheet(SHEET)
        old = []
    redo = set(args)
    unknown = redo - set(verbs)
    if unknown:
        sys.exit("Nie ma takich czasowników w arkuszu Słownik: " + ", ".join(sorted(unknown)))
    have = {r[0] for r in old} - redo
    rows = [r for r in old if r[0] in have]
    added = [v for v in verbs if v not in have]
    for v in added:
        for (k, name, _), forms in zip(TENSES, conjugate(v).values()):
            rows.append([v, name] + forms)
    order = {name: i for i, (_, name, _) in enumerate(TENSES)}
    rows.sort(key=lambda r: (r[0], order.get(r[1], 99)))
    ws.delete_rows(1, ws.max_row)
    ws.append(["Czasownik", "Czas"] + PERSONS)
    for r in rows:
        ws.append([x or None for x in r[:8]])
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:H{len(rows) + 1}"
    for col, w in zip("ABCDEFGH", [22, 24, 26, 26, 26, 26, 26, 26]):
        ws.column_dimensions[col].width = w
    wb.save(SRC)
    print(f"Arkusz {SHEET}: {len(added)} czasowników dopisanych/wygenerowanych, razem {len({r[0] for r in rows})}.")


if __name__ == "__main__":
    main(sys.argv[1:])
